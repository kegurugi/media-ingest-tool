"""Spreadsheet export: _LOGS/[DRIVE]_media_log.xlsx, rebuilt from the SQLite log.

A read-only view. It is regenerated in full after every offload, written to a
temp file and atomically swapped in. Edits made in it are never read back.

Card names, reels, timecodes, checksums and dates are text cells, so Excel
can't strip leading zeros, turn timecode into clock time or put a hash into
scientific notation. Durations are real time values shown as HH:MM:SS.
"""
from __future__ import annotations

import contextlib
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fileutil import _fsync_dir
from .layout import require_mounted, xlsx_log_path
from .masterlog import connect

TEXT, GB, DURATION, INT = "text", "gb", "duration", "int"

_TITLE_FONT = Font(bold=True, size=13)
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E7E9ED")
_FAIL_FILL = PatternFill("solid", fgColor="F8C9C4")
_SUPERSEDED_FONT = Font(color="9AA0A6", italic=True)

# (header, column type, value from row)
Column = Tuple[str, str, Callable[[Any], Any]]

CLIP_COLUMNS: List[Column] = [
    ("Card", TEXT, lambda r: r["card"]),
    ("Shoot date", TEXT, lambda r: r["shoot_date"]),
    ("Camera", TEXT, lambda r: r["camera_id"]),
    ("Reel", TEXT, lambda r: r["card"][-2:]),
    ("Clip", TEXT, lambda r: r["filename"]),
    ("Path", TEXT, lambda r: r["rel_path"]),
    ("Type", TEXT, lambda r: r["kind"]),
    ("Start TC", TEXT, lambda r: r["start_timecode"]),
    ("End TC", TEXT, lambda r: r["end_timecode"]),
    ("Duration", DURATION, lambda r: r["duration"]),
    ("Codec", TEXT, lambda r: r["codec"]),
    ("Resolution", TEXT, lambda r: r["resolution"]),
    ("FPS", TEXT, lambda r: r["frame_rate"]),
    ("Created", TEXT, lambda r: (r["created"] or "").replace("T", " ")),
    ("Camera model", TEXT, lambda r: r["camera_model"]),
    ("Size (GB)", GB, lambda r: r["size"]),
    ("xxHash64", TEXT, lambda r: r["xxh64"]),
    ("Verified", TEXT, lambda r: r["verification"]),
    ("Volume", TEXT, lambda r: r["volume"]),
    ("Offloaded", TEXT, lambda r: (r["offload_started"] or "").replace("T", " ")),
    ("Notes", TEXT, lambda r: r["notes"]),
]

CARD_COLUMNS: List[Column] = [
    ("Card", TEXT, lambda r: r["card"]),
    ("Shoot date", TEXT, lambda r: r["shoot_date"]),
    ("Camera", TEXT, lambda r: r["camera_id"]),
    ("Camera model", TEXT, lambda r: r["camera_model"]),
    ("Description", TEXT, lambda r: r["camera_description"]),
    ("Reel", TEXT, lambda r: r["reel"]),
    ("Media", TEXT, lambda r: r["media_type"]),
    ("Clips", INT, lambda r: r["clip_count"]),
    ("Files", INT, lambda r: r["file_count"]),
    ("Size (GB)", GB, lambda r: r["total_bytes"]),
    ("Duration", DURATION, lambda r: r["total_duration"]),
    ("Offload started", TEXT, lambda r: (r["offload_started"] or "").replace("T", " ")),
    ("Offload finished", TEXT, lambda r: (r["offload_finished"] or "").replace("T", " ")),
    ("Operator", TEXT, lambda r: r["operator"]),
    ("Source volume", TEXT, lambda r: r["source_volume"]),
    ("Destinations", TEXT, lambda r: r["destinations"]),
    ("Result", TEXT, lambda r: r["result"]),
    ("This drive", TEXT, lambda r: r["drive_result"]),
    ("Status", TEXT, lambda r: "superseded (recopied)" if r["superseded"] else "current"),
    ("Notes", TEXT, lambda r: r["notes"]),
    ("Errors", TEXT, lambda r: r["errors"]),
]

DAY_COLUMNS: List[Column] = [
    ("Shoot date", TEXT, lambda r: r["shoot_date"]),
    ("Cards", INT, lambda r: r["cards"]),
    ("Clips", INT, lambda r: r["clips"]),
    ("GB that day", GB, lambda r: r["bytes"]),
    ("Running total GB on drive", GB, lambda r: r["running"]),
    ("Failed cards", INT, lambda r: r["failed"]),
]


def export_xlsx(drive: Path, now: Optional[datetime] = None) -> Path:
    """Regenerate the drive's spreadsheet from its SQLite log. Returns its path."""
    drive = require_mounted(drive)
    now = now or datetime.now()
    with contextlib.closing(connect(drive)) as conn:
        offloads = conn.execute("SELECT * FROM offloads ORDER BY shoot_date, offload_started, id").fetchall()
        superseded = {r["id"] for r in offloads if r["superseded"]}
        failed_offloads = {r["id"] for r in offloads if r["drive_result"] != "PASS"}
        clips = conn.execute(
            """SELECT c.* FROM clips c JOIN offloads o ON o.id = c.offload_id
               ORDER BY c.shoot_date, c.card, o.offload_started, o.id, c.rel_path"""
        ).fetchall()
        days = _days(conn)
        show = conn.execute("SELECT show_code, project_name FROM offloads ORDER BY offload_started DESC LIMIT 1").fetchone()

    show_text = " ".join(p for p in (show["show_code"], show["project_name"]) if p) if show else ""
    title = f"{drive.name}  ·  {show_text}  ·  exported {now:%Y-%m-%d %H:%M}  ·  read-only export: edits are not kept"

    wb = Workbook()
    _sheet(wb.active, "Clips", title, CLIP_COLUMNS, clips,
           failed=lambda r: r["verification"] != "OK", grey=lambda r: r["offload_id"] in superseded)
    _sheet(wb.create_sheet(), "Cards", title, CARD_COLUMNS, offloads,
           failed=lambda r: r["id"] in failed_offloads, grey=lambda r: r["id"] in superseded)
    _sheet(wb.create_sheet(), "Days", title, DAY_COLUMNS, days,
           failed=lambda r: r["failed"] > 0, grey=lambda r: False)

    target = xlsx_log_path(drive)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=".xlsx.tmp", dir=target.parent)
    os.close(fd)
    try:
        wb.save(tmp_name)
        with open(tmp_name, "rb+") as f:
            os.fsync(f.fileno())
        os.chmod(tmp_name, 0o644)
        os.replace(tmp_name, target)
    except BaseException:
        with contextlib.suppress(FileNotFoundError):
            os.unlink(tmp_name)  # our own temp file
        raise
    _fsync_dir(target.parent)
    return target


def _days(conn) -> List[dict]:
    """Per shoot date, current (non-superseded) offloads only, with a running total."""
    rows = conn.execute(
        """SELECT shoot_date, COUNT(*) AS cards, SUM(clip_count) AS clips, SUM(total_bytes) AS bytes,
                  SUM(CASE WHEN drive_result = 'PASS' THEN 0 ELSE 1 END) AS failed
           FROM offloads WHERE superseded = 0 GROUP BY shoot_date ORDER BY shoot_date"""
    ).fetchall()
    running = 0
    days = []
    for r in rows:
        running += r["bytes"] or 0
        days.append({**dict(r), "running": running})
    return days


def _sheet(ws, name: str, title: str, columns: Sequence[Column], rows, failed, grey) -> None:
    ws.title = name
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    for col, (header, _, _) in enumerate(columns, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(h) + 2 for h, _, _ in columns]

    for i, row in enumerate(rows, 3):
        is_failed, is_grey = failed(row), grey(row)
        for col, (_, kind, get) in enumerate(columns, 1):
            value = get(row)
            cell = ws.cell(row=i, column=col)
            if kind == TEXT:
                cell.value = None if value in (None, "") else str(value)  # blank stays truly empty
                cell.number_format = "@"
                shown = cell.value or ""
            elif kind == GB:
                cell.value = round((value or 0) / 1e9, 6)
                cell.number_format = "0.00"
                shown = f"{cell.value:.2f}"
            elif kind == DURATION:
                cell.value = None if value is None else value / 86400
                cell.number_format = "[h]:mm:ss"
                shown = "00:00:00"
            else:
                cell.value = value
                shown = str(value)
            if is_failed:
                cell.fill = _FAIL_FILL
            if is_grey:
                cell.font = _SUPERSEDED_FONT
            widths[col - 1] = max(widths[col - 1], min(len(shown) + 2, 60))

    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"
    last_row = max(2, ws.max_row)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{last_row}"
