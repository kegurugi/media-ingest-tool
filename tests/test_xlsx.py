import os
from datetime import datetime

import pytest
from openpyxl import load_workbook

from conftest import SMALL_CHUNK, write_card
from media_ingest import offload
from media_ingest.naming import CardName
from media_ingest.offload import JobInfo, run_offload
from media_ingest.xlsx_export import export_xlsx


def offload_card(source, drives, name):
    card = CardName.parse(name)
    return run_offload(source, drives, card, "VIDEO", allow_same_volume=True, chunk_size=SMALL_CHUNK,
                       job=JobInfo(project_name="Example Show", show_code="ABC", camera_id=card.cam,
                                   camera_model="Sony FX6", operator="DIT", notes="note"))


@pytest.fixture
def workbook(sample_card, drives, tmp_path, monkeypatch):
    other = write_card(tmp_path / "B001", {"CLIP/C0001.MP4": b"b" * 2_000_000})
    offload_card(sample_card, drives, "ABC261012A01")
    real = offload.copy_card

    def corrupt(listing, card_dirs, **kw):
        r = real(listing, card_dirs, **kw)
        (card_dirs[0] / "CLIP/C0001.MP4").write_bytes(b"x" * 2_000_000)
        return r

    monkeypatch.setattr(offload, "copy_card", corrupt)
    offload_card(other, drives, "ABC261012B01")  # fails on drive A
    monkeypatch.setattr(offload, "copy_card", real)
    offload_card(other, drives, "ABC261013B01")
    path = export_xlsx(drives[0], now=datetime(2026, 10, 13, 9, 30))
    return load_workbook(path)


def header(ws):
    return [c.value for c in ws[2]]


def column(ws, name):
    idx = header(ws).index(name)
    return [row[idx] for row in ws.iter_rows(min_row=3)]


def test_sheets_and_title(workbook):
    assert workbook.sheetnames == ["Clips", "Cards", "Days"]
    for ws in workbook:
        title = ws["A1"].value
        assert "SHOW01_DRV01A" in title and "ABC Example Show" in title and "exported 2026-10-13 09:30" in title
        assert "edits are not kept" in title
        assert ws.freeze_panes == "A3"
        assert ws.auto_filter.ref.startswith("A2:")
        assert all(c.font.bold for c in ws[2])


def test_text_cells_keep_their_form(workbook):
    ws = workbook["Clips"]
    for name in ("Card", "Reel", "Start TC", "End TC", "xxHash64", "Shoot date", "Created"):
        for cell in column(ws, name):
            assert cell.number_format == "@", name
            assert cell.value is None or cell.data_type == "s", name
    assert [c.value for c in column(ws, "Reel")] == ["01", "01", "01", "01"]
    assert column(ws, "Shoot date")[0].value == "2026-10-12"
    hashes = [c.value for c in column(ws, "xxHash64")]
    assert all(isinstance(h, str) and len(h) == 16 for h in hashes)


def test_numbers_and_durations(workbook):
    cards = workbook["Cards"]
    size = column(cards, "Size (GB)")[0]
    assert size.number_format == "0.00" and isinstance(size.value, float)
    assert column(cards, "Duration")[0].number_format == "[h]:mm:ss"


def test_failed_rows_are_red(workbook):
    clips = workbook["Clips"]
    verified = column(clips, "Verified")
    failed = [c for c in verified if c.value != "OK"]
    assert len(failed) == 1 and failed[0].fill.fgColor.rgb.endswith("F8C9C4")
    assert all(c.fill.fill_type is None for c in verified if c.value == "OK")
    cards = workbook["Cards"]
    results = {c.value: c for c in column(cards, "Card")}
    assert results["ABC261012B01"].fill.fgColor.rgb.endswith("F8C9C4")
    assert results["ABC261012A01"].fill.fill_type is None


def test_days_with_running_total(workbook):
    days = workbook["Days"]
    values = [[c.value for c in row] for row in days.iter_rows(min_row=3)]
    assert [v[0] for v in values] == ["2026-10-12", "2026-10-13"]
    assert [v[1] for v in values] == [2, 1]
    assert values[1][4] == pytest.approx(values[0][3] + values[1][3])
    assert values[0][5] == 1  # one failed card on 10/12 on this drive


def test_export_replaces_atomically(sample_card, drives, monkeypatch):
    offload_card(sample_card, drives, "ABC261012A01")
    path = drives[0] / "_LOGS/SHOW01_DRV01A_media_log.xlsx"
    before = path.read_bytes()

    def crash(src, dst):
        raise OSError("crash")

    from media_ingest import xlsx_export
    monkeypatch.setattr(xlsx_export.os, "replace", crash)
    with pytest.raises(OSError):
        export_xlsx(drives[0])
    assert path.read_bytes() == before
    assert not [p for p in path.parent.iterdir() if p.name.endswith(".tmp")]
