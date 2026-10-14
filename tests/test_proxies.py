"""Proxies: ffmpeg-generated clips, checked with ffprobe and by reading pixels."""
import json
import shutil
import subprocess
from pathlib import Path

import pytest

from conftest import SMALL_CHUNK
from media_ingest import offload as offload_mod
from media_ingest import proxies
from media_ingest.cli import main
from media_ingest.config import config_from_dict
from media_ingest.errors import ConfigError
from media_ingest.naming import CardName
from media_ingest.offload import JobInfo, run_offload

pytestmark = pytest.mark.skipif(not shutil.which("ffmpeg"), reason="needs ffmpeg")
CARD = CardName.parse("ABC261012A01")
EXAMPLE = json.loads((Path(__file__).resolve().parent.parent / "project.example.json").read_text())
NO_BLUE = "LUT_3D_SIZE 2\n0 0 0\n1 0 0\n0 1 0\n1 1 0\n0 0 0\n1 0 0\n0 1 0\n1 1 0\n"
SONY_XML = ('<?xml version="1.0" encoding="UTF-8"?>\n<NonRealTimeMeta xmlns="urn:schemas-professionalDisc:'
            'nonRealTimeMeta:ver.2.20"><CreationDate value="2026-10-12T09:00:00-08:00"/>'
            '<AcquisitionRecord><Group name="CameraUnitMetadataSet"><Item name="CaptureGammaEquation" '
            'value="{gamma}"/></Group></AcquisitionRecord></NonRealTimeMeta>\n')


def ffmpeg(*args):
    subprocess.run(["ffmpeg", "-v", "error", "-y", *map(str, args)], check=True)


def probe(path):
    out = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams",
                          str(path)], capture_output=True, text=True).stdout
    return json.loads(out)


def mean_rgb(path):
    raw = subprocess.run(["ffmpeg", "-v", "error", "-i", str(path), "-frames:v", "1", "-vf", "scale=32:18",
                          "-f", "rawvideo", "-pix_fmt", "rgb24", "-"], capture_output=True).stdout
    n = len(raw) // 3
    return [sum(raw[i::3]) / n for i in range(3)]


def make_card(root, gamma=None):
    """A Sony-style card: a UHD MXF with two audio tracks (+ XML), and a small MP4."""
    mxf = root / "XDROOT/Clip/A002C001_26101201.MXF"
    mxf.parent.mkdir(parents=True)
    ffmpeg("-f", "lavfi", "-i", "testsrc2=size=3840x2160:rate=24000/1001", "-f", "lavfi", "-i",
           "sine=f=440:sample_rate=48000", "-f", "lavfi", "-i", "sine=f=880:sample_rate=48000", "-t", 1,
           "-map", 0, "-map", 1, "-map", 2, "-c:v", "mpeg2video", "-q:v", 8, "-c:a", "pcm_s16le",
           "-timecode", "10:38:49:16", "-f", "mxf", mxf)
    if gamma:
        mxf.with_name("A002C001_26101201M01.XML").write_text(SONY_XML.format(gamma=gamma))
    mp4 = root / "XDROOT/Clip/C0002.MP4"
    ffmpeg("-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25", "-t", 1, "-c:v", "libx264",
           "-timecode", "01:00:00:00", mp4)
    return root


def config(tmp_path, **proxy):
    data = json.loads(json.dumps(EXAMPLE))
    if proxy:
        data["proxies"] = proxy
    cfg = config_from_dict(data)
    cfg.base_dir = tmp_path
    return cfg


@pytest.fixture
def offloaded(tmp_path, drives):
    def go(gamma=None, corrupt=None):
        src = make_card(tmp_path / "card", gamma)
        r = run_offload(src, drives, CARD, "VIDEO", allow_same_volume=True, chunk_size=SMALL_CHUNK,
                        job=JobInfo(show_code="ABC", camera_id="A"))
        assert r.passed
        return r
    return go


def test_prores_proxies_on_both_drives(offloaded, drives, tmp_path):
    offloaded()
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path), drives[1:])
    assert out.ok and out.copied_to == ["SHOW01_DRV01B"], out
    for d in drives:
        folder = d / "10122026/PROXIES/ABC261012A01"
        assert sorted(p.name for p in folder.iterdir() if not p.name.startswith(".")) == [
            "A002C001_26101201.mov", "C0002.mov"]
    uhd = probe(drives[1] / "10122026/PROXIES/ABC261012A01/A002C001_26101201.mov")
    video = next(s for s in uhd["streams"] if s["codec_type"] == "video")
    assert (video["codec_name"], video["profile"], video["width"], video["height"]) == ("prores", "Proxy", 1920, 1080)
    audio = [s for s in uhd["streams"] if s["codec_type"] == "audio"]
    assert len(audio) == 2
    tmcd = next(s for s in uhd["streams"] if s.get("tags", {}).get("reel_name"))
    assert tmcd["tags"] == {**tmcd["tags"], "timecode": "10:38:49:16", "reel_name": "ABC261012A01"}
    assert abs(float(uhd["format"]["duration"]) - 1.0) < 0.1
    small = probe(drives[0] / "10122026/PROXIES/ABC261012A01/C0002.mov")
    video = next(s for s in small["streams"] if s["codec_type"] == "video")
    assert (video["width"], video["height"]) == (640, 360)  # never upscaled
    a = drives[0] / "10122026/PROXIES/ABC261012A01/C0002.mov"
    b = drives[1] / "10122026/PROXIES/ABC261012A01/C0002.mov"
    assert a.read_bytes() == b.read_bytes()


def test_results_logged_and_in_spreadsheet(offloaded, drives, tmp_path):
    import contextlib
    from openpyxl import load_workbook
    from media_ingest import masterlog
    offloaded()
    proxies.run_for_card(drives[0], CARD, config(tmp_path), drives[1:])
    for d in drives:
        with contextlib.closing(masterlog.connect(d)) as conn:
            rows = [dict(r) for r in conn.execute("SELECT * FROM proxies ORDER BY clip")]
        assert [(r["clip"].rsplit("/", 1)[1], r["result"], r["volume"]) for r in rows] == [
            ("A002C001_26101201.MXF", "OK", d.name), ("C0002.MP4", "OK", d.name)]
        assert rows[0]["start_timecode"] == "10:38:49:16" and len(rows[0]["xxh64"]) == 16
        assert "Proxies" in load_workbook(d / f"_LOGS/{d.name}_media_log.xlsx").sheetnames


def test_h264_option(offloaded, drives, tmp_path):
    offloaded()
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path, codec="h264"), drives[1:])
    assert out.ok
    info = probe(drives[0] / "10122026/PROXIES/ABC261012A01/A002C001_26101201.mov")
    assert next(s for s in info["streams"] if s["codec_type"] == "video")["codec_name"] == "h264"
    assert any(s.get("tags", {}).get("reel_name") == "ABC261012A01" for s in info["streams"])


def test_lut_is_burned_in(offloaded, drives, tmp_path):
    offloaded()
    (tmp_path / "noblue.cube").write_text(NO_BLUE)
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path, lut="noblue.cube"), drives[1:])
    assert out.ok and all(r.lut == "noblue.cube" for r in out.results)
    r, g, b = mean_rgb(drives[1] / "10122026/PROXIES/ABC261012A01/A002C001_26101201.mov")
    assert b < 6 and r > 40 and g > 40
    video = next(s for s in probe(drives[0] / "10122026/PROXIES/ABC261012A01/C0002.mov")["streams"]
                 if s["codec_type"] == "video")
    assert video.get("color_transfer") == "bt709"


@pytest.mark.parametrize("gamma, applied", [("rec709", False), ("s-log3-cine", True), ("S-Log3", True)])
def test_lut_respects_sony_capture_gamma(offloaded, drives, tmp_path, gamma, applied):
    offloaded(gamma=gamma)
    (tmp_path / "luts").mkdir()
    (tmp_path / "luts/slog3.cube").write_text(NO_BLUE)
    data = json.loads(json.dumps(EXAMPLE))
    data["cameras"]["A"] = {"model": "Sony FX6", "lut": "luts/slog3.cube", "lut_gamma": "s-log3"}
    cfg = config_from_dict(data)
    cfg.base_dir = tmp_path
    out = proxies.run_for_card(drives[0], CARD, cfg, drives[1:])
    sony = next(r for r in out.results if r.clip.endswith(".MXF"))
    blue = mean_rgb(sony.proxy)[2]
    if applied:
        assert sony.lut == "slog3.cube" and blue < 6
    else:
        assert sony.lut == "" and blue > 40 and "clip is rec709, LUT is for s-log3: LUT not applied" == sony.lut_note
    mp4 = next(r for r in out.results if r.clip.endswith(".MP4"))  # no XML: gamma unknown, LUT applied
    assert mp4.lut == "slog3.cube" and "gamma unknown" in mp4.lut_note


def test_missing_lut_makes_no_proxies(offloaded, drives, tmp_path):
    offloaded()
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path, lut="nope.cube"), drives[1:])
    assert "LUT problem" in out.problem and not (drives[0] / "10122026/PROXIES").exists()


def test_failed_card_gets_no_proxies(tmp_path, drives, monkeypatch):
    src = make_card(tmp_path / "card")
    real = offload_mod.copy_card
    monkeypatch.setattr(offload_mod, "copy_card",
                        lambda l, d, **kw: (real(l, d, **kw), (d[1] / "XDROOT/Clip/C0002.MP4").write_bytes(b"x"))[0])
    run_offload(src, drives, CARD, "VIDEO", allow_same_volume=True, job=JobInfo(show_code="ABC", camera_id="A"))
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path), drives[1:])
    assert "no proxies for unverified media" in out.problem


def test_rerun_never_overwrites_and_fills_in_missing_copies(offloaded, drives, tmp_path):
    offloaded()
    proxies.run_for_card(drives[0], CARD, config(tmp_path), [])  # second drive "not connected"
    first = drives[0] / "10122026/PROXIES/ABC261012A01/C0002.mov"
    before = first.stat().st_mtime_ns
    assert not (drives[1] / "10122026/PROXIES").exists()
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path), drives[1:])
    assert {r.status for r in out.results} == {"SKIPPED"} and out.copied_to == ["SHOW01_DRV01B"]
    assert first.stat().st_mtime_ns == before
    assert (drives[1] / "10122026/PROXIES/ABC261012A01/C0002.mov").read_bytes() == first.read_bytes()


def test_disconnected_drive_warns(offloaded, drives, tmp_path):
    offloaded()
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path), [tmp_path / "GONE"])
    assert out.ok and any("isn't connected" in w for w in out.warnings)


def test_software_fallback(offloaded, drives, tmp_path):
    offloaded()
    out = proxies.make_proxies(drives[0], CARD, config(tmp_path), [], encoders={"prores_ks", "libx264"})
    assert out.ok and {r.codec for r in out.results} == {"prores_ks"}


def test_bad_clip_fails_without_leftovers(offloaded, drives, tmp_path):
    offloaded()
    from media_ingest.records import read_card_record
    record = read_card_record(drives[0] / "_REPORTS/10122026/ABC261012A01_card.json")
    clip = drives[0] / "10122026/VIDEO/ABC261012A01/XDROOT/Clip/C0002.MP4"
    clip.chmod(0o644)
    clip.write_bytes(b"\0" * clip.stat().st_size)  # unreadable as video
    out = proxies.run_for_card(drives[0], CARD, config(tmp_path), [])
    bad = next(r for r in out.results if r.clip.endswith("C0002.MP4"))
    assert bad.status == "FAIL" and not out.ok
    folder = drives[0] / "10122026/PROXIES/ABC261012A01"
    assert sorted(p.name for p in folder.iterdir()) == ["A002C001_26101201.mov"]


def test_off_by_default_and_auto(tmp_path, drives, config_file, monkeypatch, capsys):
    from conftest import save_day_details
    save_day_details(config_file, "2026-10-12")
    src = make_card(tmp_path / "card")
    argv = ["--config", str(config_file), "ingest", "--source", str(src), "--dest", str(drives[0]),
            "--dest", str(drives[1]), "--camera", "A", "--date", "10122026", "--allow-same-volume"]
    answers = iter(["", "", "", "y"])
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    assert main(argv) == 0
    assert not (drives[0] / "10122026/PROXIES").exists()  # off by default: nothing at all
    data = json.loads(config_file.read_text())
    data["proxies"] = {"auto": True}
    config_file.write_text(json.dumps(data))
    answers = iter(["", "", "", "y"])
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    assert main(argv) == 0
    out = capsys.readouterr().out
    assert "Making proxies" in out
    assert (drives[1] / "10122026/PROXIES/ABC261012A02/C0002.mov").exists()


def test_mit_proxies_command(offloaded, drives, config_file, capsys):
    offloaded()
    assert main(["--config", str(config_file), "proxies", "--drive", str(drives[0])]) == 0
    out = capsys.readouterr().out
    assert "OK       C0002.MP4 -> C0002.mov" in out and "Copied and verified on: SHOW01_DRV01B" in out


@pytest.mark.parametrize("proxy, message", [
    ({"codec": "dnxhd"}, "proxies.codec"), ({"max_width": 1921}, "max_width"),
    ({"auto": "yes"}, "proxies.auto"), ({"burn": True}, "unknown key"), ({"lut": ""}, "proxies.lut"),
])
def test_proxy_settings_validation(proxy, message):
    data = json.loads(json.dumps(EXAMPLE))
    data["proxies"] = proxy
    with pytest.raises(ConfigError, match=message):
        config_from_dict(data)
