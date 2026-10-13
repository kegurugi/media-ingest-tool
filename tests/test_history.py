"""Project history across drive sets: running totals and reel numbers carry over to new drives."""
import contextlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import SMALL_CHUNK, write_card
from media_ingest import history, masterlog
from media_ingest.cli import main
from media_ingest.history import history_path, summary, used_reels
from media_ingest.naming import CardName
from media_ingest.offload import JobInfo, run_offload
from media_ingest.reels import propose_reel

OCT12 = date(2026, 10, 12)


@pytest.fixture
def kv(tmp_path):
    """Two drive sets: DRV01A/B (first) and DRV02A/B (continued on)."""
    sets = {}
    for name in ("DRV01", "DRV02"):
        pair = [tmp_path / f"{name}A", tmp_path / f"{name}B"]
        for d in pair:
            d.mkdir()
        sets[name] = pair
    return sets


@pytest.fixture
def hist(tmp_path):
    return tmp_path / "project" / "project_history.sqlite"


def offload(source, drives, name, hist):
    card = CardName.parse(name)
    return run_offload(source, drives, card, "VIDEO", allow_same_volume=True, chunk_size=SMALL_CHUNK,
                       job=JobInfo(show_code="ABC", camera_id=card.cam), project_history=hist)


@pytest.fixture
def cards(tmp_path):
    return [write_card(tmp_path / f"card{i}", {"CLIP/C0001.MP4": bytes([i]) * (1_000_000 * (i + 1))}) for i in range(4)]


def test_history_path_is_beside_project_json(tmp_path):
    assert history_path(tmp_path / "proj" / "project.json") == (tmp_path / "proj" / "project_history.sqlite").resolve()


def test_totals_continue_on_the_next_drive_set(kv, hist, cards):
    assert offload(cards[0], kv["DRV01"], "ABC261012A01", hist).passed
    assert offload(cards[1], kv["DRV01"], "ABC261012A02", hist).passed
    # DRV01 is full; the day continues on DRV02.
    assert offload(cards[2], kv["DRV02"], "ABC261012A03", hist).passed
    rows = summary(hist)
    assert [(r["drive_set"], r["cards"], r["bytes"]) for r in rows] == [
        ("DRV01A & DRV01B", 2, 3_000_000), ("DRV02A & DRV02B", 1, 3_000_000)]

    # DRV02 carries DRV01's totals in its own log and spreadsheet.
    with contextlib.closing(masterlog.connect(kv["DRV02"][0])) as conn:
        assert {r["drive_set"] for r in conn.execute("SELECT drive_set FROM project_summary")} == {
            "DRV01A & DRV01B", "DRV02A & DRV02B"}
    wb = load_workbook(kv["DRV02"][0] / "_LOGS/DRV02A_media_log.xlsx")
    ws = wb["All drives"]
    h = [c.value for c in ws[2]]
    data = [dict(zip(h, (c.value for c in row))) for row in ws.iter_rows(min_row=3, max_row=4)]
    assert [(r["Drive set"], r["This drive"], r["Cards"]) for r in data] == [
        ("DRV01A & DRV01B", None, 2), ("DRV02A & DRV02B", "yes", 1)]
    assert data[1]["Project running total GB"] == pytest.approx(0.006)
    days = wb["Days"]
    h = [c.value for c in days[2]]
    day = dict(zip(h, (c.value for c in days[3])))
    assert day["Running total GB on drive"] == pytest.approx(0.003)  # just this drive
    assert day["Project total GB (all drives)"] == pytest.approx(0.006)  # everything so far


def test_reels_never_reused_across_drive_sets(kv, hist, cards):
    offload(cards[0], kv["DRV01"], "ABC261012A01", hist)
    offload(cards[1], kv["DRV01"], "ABC261012A02", hist)
    assert propose_reel(kv["DRV02"], "ABC", OCT12, "A") == 1  # the drives alone don't know
    elsewhere = used_reels(hist, "ABC", OCT12, "A")
    assert elsewhere == {1, 2}
    assert propose_reel(kv["DRV02"], "ABC", OCT12, "A", also_used=elsewhere) == 3
    assert used_reels(hist, "ABC", OCT12, "B") == set()


def test_superseded_attempts_not_counted(kv, hist, cards, monkeypatch):
    from media_ingest import offload as offload_mod
    from media_ingest.incomplete import find_incomplete_all, quarantine
    real = offload_mod.copy_card

    def corrupt(listing, card_dirs, **kw):
        r = real(listing, card_dirs, **kw)
        (card_dirs[0] / "CLIP/C0001.MP4").write_bytes(b"x")
        return r

    monkeypatch.setattr(offload_mod, "copy_card", corrupt)
    assert not offload(cards[0], kv["DRV01"], "ABC261012A01", hist).passed
    monkeypatch.setattr(offload_mod, "copy_card", real)
    assert summary(hist)[0]["failed"] == 1
    for item, d in zip(find_incomplete_all(kv["DRV01"], "ABC", OCT12, "A")["ABC261012A01"], kv["DRV01"]):
        quarantine(item, d)
    assert offload(cards[0], kv["DRV01"], "ABC261012A01", hist).passed
    (row,) = summary(hist)
    assert (row["cards"], row["failed"], row["bytes"]) == (1, 0, 1_000_000)


def test_rebuild_from_drives(kv, hist, cards, tmp_path):
    offload(cards[0], kv["DRV01"], "ABC261012A01", hist)
    offload(cards[1], kv["DRV02"], "ABC261012A02", hist)
    before = summary(hist)
    hist.rename(tmp_path / "lost.sqlite")
    added, _ = history.rebuild_from_drives(hist, [kv["DRV01"][0], kv["DRV02"][0]])
    assert added == 2 and summary(hist) == before
    assert history.rebuild_from_drives(hist, [kv["DRV01"][0]])[0] == 0


def test_damaged_history_is_set_aside_on_rebuild(kv, hist, cards):
    offload(cards[0], kv["DRV01"], "ABC261012A01", hist)
    hist.write_bytes(b"garbage" * 100)
    r = offload(cards[1], kv["DRV01"], "ABC261012A02", hist)
    assert r.passed and any("Project history" in w for w in r.warnings)
    added, aside = history.rebuild_from_drives(hist, [kv["DRV01"][0]])
    assert added == 2 and aside.read_bytes().startswith(b"garbage")


def test_no_history_still_works(kv, cards):
    r = run_offload(cards[0], kv["DRV01"], CardName.parse("ABC261012A01"), "VIDEO", allow_same_volume=True)
    assert r.passed and not r.warnings
    wb = load_workbook(kv["DRV01"][0] / "_LOGS/DRV01A_media_log.xlsx")
    assert "No project history" in str([c.value for row in wb["All drives"].iter_rows() for c in row])


def test_ingest_uses_history_for_reels(kv, cards, tmp_path, monkeypatch, capsys):
    config = tmp_path / "project" / "project.json"
    config.parent.mkdir()
    config.write_text((Path(__file__).resolve().parent.parent / "project.example.json").read_text())
    hist = history_path(config)
    offload(cards[0], kv["DRV01"], "ABC261012A01", hist)
    answers = iter(["", "", "", "y"])
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    code = main(["--config", str(config), "ingest", "--source", str(cards[1]), "--dest", str(kv["DRV02"][0]),
                 "--dest", str(kv["DRV02"][1]), "--camera", "A", "--date", "10122026", "--allow-same-volume"])
    out = capsys.readouterr().out
    assert code == 0, out
    assert "also has reel(s) 01 on 10/12/2026 in the project history" in out
    assert (kv["DRV02"][0] / "10122026/VIDEO/ABC261012A02").is_dir()


def test_history_commands(kv, cards, tmp_path, capsys):
    config = tmp_path / "project.json"
    config.write_text((Path(__file__).resolve().parent.parent / "project.example.json").read_text())
    offload(cards[0], kv["DRV01"], "ABC261012A01", None)  # offloaded without a history
    assert main(["--config", str(config), "history", "show"]) == 2
    assert main(["--config", str(config), "history", "rebuild", "--dest", str(kv["DRV01"][0]),
                 "--dest", str(kv["DRV01"][1])]) == 0
    out = capsys.readouterr().out
    assert "1 card record(s) added" in out and "DRV01A & DRV01B: 1 card(s)" in out and "Project total: 1 card(s)" in out
    with contextlib.closing(masterlog.connect(kv["DRV01"][1])) as conn:
        assert conn.execute("SELECT COUNT(*) FROM project_summary").fetchone()[0] == 1
