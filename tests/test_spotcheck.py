"""`mit verify`: re-hashing card folders against the log."""
import contextlib
import os
import sqlite3

import pytest
from openpyxl import load_workbook

from conftest import SMALL_CHUNK, write_card
from media_ingest import masterlog
from media_ingest.cli import main
from media_ingest.errors import IngestError
from media_ingest.incomplete import find_incomplete_all, quarantine
from media_ingest.naming import CardName
from media_ingest.offload import JobInfo, run_offload
from media_ingest import offload as offload_mod
from media_ingest.spotcheck import check_card, locate_card

CLIP1 = "PRIVATE/M4ROOT/CLIP/C0001.MP4"


def offload(source, drives, name="ABC261012A01"):
    card = CardName.parse(name)
    r = run_offload(source, drives, card, "VIDEO", allow_same_volume=True, chunk_size=SMALL_CHUNK,
                    job=JobInfo(show_code="ABC", camera_id=card.cam))
    return r


@pytest.fixture
def offloaded(sample_card, drives):
    assert offload(sample_card, drives).passed
    return [d / "10122026/VIDEO/ABC261012A01" for d in drives]


def verify(*args):
    return main(["verify", *map(str, args)])


def flip_byte(path, offset=1000):
    data = bytearray(path.read_bytes())
    data[offset] ^= 0x01
    path.write_bytes(bytes(data))


def verifications(drive):
    with contextlib.closing(masterlog.connect(drive)) as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM verifications ORDER BY id")]


def test_clean_card_passes(offloaded, capsys):
    assert verify(offloaded[1]) == 0
    out = capsys.readouterr().out
    assert "PASS  ABC261012A01  5 files" in out and "hashes from the master log" in out
    assert "PASS  1 card(s) verified" in out


def test_one_changed_byte_fails(offloaded, capsys):
    """The spec's corruption test."""
    flip_byte(offloaded[1] / CLIP1)
    assert verify(offloaded[1]) == 1
    captured = capsys.readouterr()
    assert "FAIL  ABC261012A01" in captured.out
    assert "VERIFY FAILED: 1 of 1 card(s)" in captured.err and f"MISMATCH {CLIP1}" in captured.err
    assert verify(offloaded[0]) == 0  # the other drive's copy is still good


def test_missing_truncated_and_missing_folder(offloaded):
    folder = offloaded[0]
    os.rename(folder / "PRIVATE/M4ROOT/MEDIAPRO.XML", folder.parent / "elsewhere.xml")
    with open(folder / "PRIVATE/M4ROOT/CLIP/C0002.MP4", "r+b") as f:
        f.truncate(10)
    (folder / "PRIVATE/M4ROOT/SUB").rmdir()
    check = check_card(folder)
    assert not check.passed
    assert {(c.rel_path, c.status) for c in check.failures} == {
        ("PRIVATE/M4ROOT/MEDIAPRO.XML", "MISSING"), ("PRIVATE/M4ROOT/CLIP/C0002.MP4", "SIZE")}
    assert check.missing_dirs == ["PRIVATE/M4ROOT/SUB"]


def test_clutter_and_report_copy_ignored_extra_files_warn(offloaded):
    folder = offloaded[0]
    assert (folder / "ABC261012A01_report.html").exists()
    (folder / ".DS_Store").write_bytes(b"finder")
    (folder / "PRIVATE/._C0001.MP4").write_bytes(b"x")
    (folder / "PRIVATE/notes.txt").write_text("someone added this")
    check = check_card(folder)
    assert check.passed
    assert check.warnings == ["unexpected file PRIVATE/notes.txt"]


def test_read_only_card_folder(offloaded):
    folder = offloaded[0]
    paths = [folder, *folder.rglob("*")]
    before = {str(p): p.stat().st_mtime_ns for p in paths}
    for p in paths:
        os.chmod(p, 0o555 if p.is_dir() else 0o444)
    try:
        assert check_card(folder).passed
        assert {str(p): p.stat().st_mtime_ns for p in paths} == before
    finally:
        for p in paths:
            os.chmod(p, 0o755 if p.is_dir() else 0o644)


def test_falls_back_to_card_record(offloaded, drives, tmp_path):
    os.rename(drives[0] / "_LOGS/media_log.sqlite", tmp_path / "away.sqlite")
    check = check_card(offloaded[0])
    assert check.passed and check.source == "card record"
    assert any("hashes came from the card record" in n for n in check.notes)
    flip_byte(offloaded[0] / CLIP1)
    assert not check_card(offloaded[0]).passed


def test_log_and_record_disagreement_is_flagged(offloaded, drives):
    with contextlib.closing(sqlite3.connect(drives[0] / "_LOGS/media_log.sqlite")) as conn:
        conn.execute("UPDATE files SET xxh64 = 'deadbeefdeadbeef' WHERE path = ?", (CLIP1,))
        conn.commit()
    check = check_card(offloaded[0])
    assert not check.passed
    assert any("LOG/RECORD DISAGREE" in p and CLIP1 in p for p in check.problems)


def test_not_a_card_folder(drives, tmp_path, capsys):
    for bad in (tmp_path, drives[0] / "10122026/VIDEO", drives[0] / "10122026/OTHER/ABC261012A01",
                drives[0] / "10132026/VIDEO/ABC261012A01"):
        with pytest.raises(IngestError, match="not a card folder"):
            locate_card(bad)
    assert verify(tmp_path) == 2
    assert "not a card folder" in capsys.readouterr().err


def test_nothing_to_verify_against(drives, capsys):
    folder = drives[0] / "10122026/VIDEO/ABC261012A09"
    folder.mkdir(parents=True)
    assert verify(folder) == 2
    assert "nothing to verify against" in capsys.readouterr().err


def test_several_folders_at_once(sample_card, drives, tmp_path, capsys):
    other = write_card(tmp_path / "B001", {"CLIP/C0001.MP4": b"b" * 5000})
    offload(sample_card, drives, "ABC261012A01")
    offload(other, drives, "ABC261012B01")
    flip_byte(drives[0] / "10122026/VIDEO/ABC261012B01/CLIP/C0001.MP4", 100)
    assert verify(drives[0] / "10122026/VIDEO/ABC261012A01", drives[0] / "10122026/VIDEO/ABC261012B01") == 1
    captured = capsys.readouterr()
    assert "[1/2] PASS  ABC261012A01" in captured.out and "[2/2] FAIL  ABC261012B01" in captured.out


def test_drive_mode_finds_missing_card_folders(sample_card, drives, tmp_path, capsys):
    other = write_card(tmp_path / "B001", {"CLIP/C0001.MP4": b"b" * 5000})
    offload(sample_card, drives, "ABC261012A01")
    offload(other, drives, "ABC261012B01")
    offload(other, drives, "ABC261013B01")
    # The B01 folder went missing in transit; someone also copied a stray card folder onto the drive.
    os.rename(drives[1] / "10122026/VIDEO/ABC261012B01", tmp_path / "lost")
    (drives[1] / "10122026/VIDEO/ABC261012C07").mkdir()
    assert verify("--drive", drives[1], "--date", "10122026") == 1
    captured = capsys.readouterr()
    assert "Verifying 2 card(s)" in captured.out  # only 10/12
    assert "10122026/VIDEO/ABC261012C07 is on the drive but not in its log" in captured.out
    assert "ABC261012B01" in captured.err and "card folder missing" in captured.err
    assert verify("--drive", drives[0]) == 0
    assert "Verifying 3 card(s)" in capsys.readouterr().out


def test_drive_mode_skips_superseded_attempts(sample_card, drives, monkeypatch, capsys):
    real = offload_mod.copy_card
    monkeypatch.setattr(offload_mod, "copy_card", lambda l, d, **kw: (real(l, d, **kw), (d[0] / CLIP1).write_bytes(b"x"))[0])
    offload(sample_card, drives)
    monkeypatch.setattr(offload_mod, "copy_card", real)
    for item, d in zip(find_incomplete_all(drives, "ABC", CardName.parse("ABC261012A01").shoot_date, "A")["ABC261012A01"], drives):
        quarantine(item, d)
    assert offload(sample_card, drives).passed
    assert verify("--drive", drives[0]) == 0
    assert "Verifying 1 card(s)" in capsys.readouterr().out


def test_failed_offload_is_noted(sample_card, drives, monkeypatch, capsys):
    real = offload_mod.copy_card
    monkeypatch.setattr(offload_mod, "copy_card", lambda l, d, **kw: (real(l, d, **kw), flip_byte(d[1] / CLIP1))[0])
    offload(sample_card, drives)
    folder = drives[0] / "10122026/VIDEO/ABC261012A01"
    check = check_card(folder)
    assert check.passed  # drive A's copy is fine
    assert any("original offload was FAIL" in n for n in check.notes)
    assert any("marked incomplete (FAILED)" in n for n in check.notes)


def test_results_recorded_in_log_and_spreadsheet(offloaded, drives, capsys):
    flip_byte(offloaded[1] / CLIP1)
    verify(offloaded[0])
    verify(offloaded[1])
    assert [(r["card"], r["result"]) for r in verifications(drives[0])] == [("ABC261012A01", "PASS")]
    (row,) = verifications(drives[1])
    assert row["result"] == "FAIL" and row["failures"] == 1 and f"MISMATCH {CLIP1}" in row["details"]
    assert row["files_checked"] == 5 and row["source"] == "master log"
    ws = load_workbook(drives[1] / "_LOGS/SHOW01_DRV01B_media_log.xlsx")["Verifications"]
    h = [c.value for c in ws[2]]
    cells = dict(zip(h, ws[3]))
    assert cells["Result"].value == "FAIL" and cells["Result"].fill.fgColor.rgb.endswith("F8C9C4")


def test_no_record_writes_nothing(offloaded, drives):
    before = {p: p.stat().st_mtime_ns for p in (drives[0] / "_LOGS").iterdir()}
    assert verify(offloaded[0], "--no-record") == 0
    assert {p: p.stat().st_mtime_ns for p in (drives[0] / "_LOGS").iterdir()} == before
    assert verifications(drives[0]) == []


def test_argument_errors(drives, capsys):
    assert verify("--date", "10122026") == 2
    assert verify() == 2
    assert verify("--drive", drives[0], drives[0] / "x") == 2
    assert verify("--drive", drives[0].parent / "GONE") == 2
    err = capsys.readouterr().err
    assert "--date needs --drive" in err and "not found" in err
